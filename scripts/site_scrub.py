"""Website scrub for one or more MH Sub-Categories -> F&T Site Scrub <subcat> <date>.xlsx

  python site_scrub.py --msl MSL.xlsx --ds DS.xlsx --shopify shopify_variants.tsv --subcat "Uph Seating" --out-dir .

Should be online = MSL rows in MH Division "Furniture + Textiles", the chosen MH Sub-Category(s),
Stage in STAGES. Each of those SKUs is matched (by SKU) to the live site (magnolia.com/products.json).
SKUs that aren't live get a Reason, checked in this order:
  Not in Shopify -> Draft/Archived in Shopify -> No inventory (<= 0) -> Under stock buffer (<= BUFFER)
  -> Has inventory - investigate
Inventory = DS report Quantity for dropship SKUs, Jewell Warehouse available for the rest.

shopify_variants.tsv = Shopify admin pull (see README step 2): header, then sku, status, available, on_hand.
"""
import argparse
import datetime
import json
import os
import time
import urllib.request

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

DIVISION = "Furniture + Textiles"
STAGES = ["Continuing", "New", "Markdown", "Residual", "Re-Intro"]
BUFFER = 10
INVESTIGATE = "Has inventory - investigate"


def read_sheet(path, need):
    """Rows as dicts from the first sheet whose header row has every column in `need`."""
    wb = openpyxl.load_workbook(path, read_only=True)
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        head = [str(h).strip() if h is not None else "" for h in next(rows, [])]
        if all(n in head for n in need):
            return [dict(zip(head, r)) for r in rows]
    raise SystemExit(f"no sheet in {path} has columns {need}")


def read_shopify(path):
    """sku -> (status, available). Prefers the ACTIVE record when a SKU is on more than one variant."""
    out = {}
    with open(path, encoding="utf-8-sig") as f:
        for line in f.read().splitlines()[1:]:
            p = (line.split("\t") + [""] * 4)[:4]
            sku, status, avail = p[0].strip(), p[1], p[2]
            if not sku or (sku in out and out[sku][0] == "ACTIVE"):
                continue
            out[sku] = (status, int(avail) if avail not in ("", None) else None)
    return out


def pull_site(log=print):
    prods, page = [], 1
    while True:
        req = urllib.request.Request(f"https://magnolia.com/products.json?limit=250&page={page}",
                                     headers={"User-Agent": "Mozilla/5.0"})
        batch = json.loads(urllib.request.urlopen(req, timeout=60).read()).get("products", [])
        if not batch:
            break
        prods += batch
        page += 1
        time.sleep(0.5)
    log(f"site: {len(prods)} products")
    return prods


def tier(img):
    s = min(img["width"], img["height"])
    return "Low" if s < 1000 else "OK" if s < 1500 else "Good"


def site_variants(prods):
    """sku -> one row per live variant, with the image a shopper sees for it."""
    out = {}
    for p in prods:
        main = p["images"][0] if p["images"] else None
        multi = len(p["variants"]) > 1
        for v in p["variants"]:
            sku = (v.get("sku") or "").strip()
            if not sku:
                continue
            img = v.get("featured_image") or (None if multi else main)
            if img:
                quality, size = tier(img), f"{img['width']}x{img['height']}"
            elif main:
                quality, size = "No variant photo", f"main {main['width']}x{main['height']} ({tier(main)})"
            else:
                quality, size = "No image", ""
            out[sku] = {"product": p["title"], "variant": "" if v["title"] == "Default Title" else v["title"],
                        "vendor": p.get("vendor") or "", "in_stock": "Yes" if v.get("available") else "No",
                        "quality": quality, "size": size,
                        "url": f"https://magnolia.com/products/{p['handle']}?variant={v['id']}"}
    return out


def reason(status, inv):
    if status is None:
        return "Not in Shopify"
    if status != "ACTIVE":
        return f"{status.title()} in Shopify"
    if inv is None or inv <= 0:
        return "No inventory"
    if inv <= BUFFER:
        return "Under stock buffer"
    return INVESTIGATE


def write_tab(wb, title, headers, rows, widths):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font = Font(name="Arial", bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="3F3F3F")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name="Arial")
        row[0].number_format = "@"
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    return ws


def build(msl_path, ds_path, shopify_path, subcats, out_dir, prods=None):
    msl = [r for r in read_sheet(msl_path, ["SKU", "Stage", "MH Division", "MH Sub-Category"])
           if r["MH Division"] == DIVISION and r["MH Sub-Category"] in subcats and r["Stage"] in STAGES and r["SKU"]]
    if not msl:
        raise SystemExit(f"no {DIVISION} SKUs in {subcats} at stages {STAGES}")
    ds = {str(r["SKU"]).strip(): r["Quantity"] for r in read_sheet(ds_path, ["SKU", "Quantity"]) if r["SKU"]} if ds_path else {}
    shop = read_shopify(shopify_path)
    site = site_variants(prods if prods is not None else pull_site())

    scrub, outliers = [], []
    for m in sorted(msl, key=lambda m: (m["MH Sub-Category"], m["Display Name"] or "")):
        sku = str(m["SKU"]).strip()
        s = site.get(sku)
        dropship = m.get("Is Dropship Item") == "Yes"
        if dropship:
            inv, src = (ds[sku], "DS") if sku in ds else (None, "DS (not on report)")
        else:
            inv = shop[sku][1] if sku in shop else None
            src = "Jewell" if inv is not None else "Jewell (no record)"
        status = shop[sku][0] if sku in shop else None
        base = [sku, m["Display Name"], s["variant"] if s else "", "Yes" if s else "No",
                s["quality"] if s else "", inv]
        extra = [src, status or "", m["MH Sub-Category"], m["Vendor"], m["Stage"], s["size"] if s else ""]
        scrub.append(base + extra)
        if not s:
            outliers.append([sku, m["Display Name"], reason(status, inv), inv, src, status or "",
                             m["MH Sub-Category"], m["Vendor"], m["Stage"]])
    order = [INVESTIGATE, "Not in Shopify", "Draft in Shopify", "Archived in Shopify", "Under stock buffer", "No inventory"]
    outliers.sort(key=lambda r: (order.index(r[2]) if r[2] in order else len(order), r[1] or ""))

    msl_skus = {str(m["SKU"]).strip() for m in msl}
    html = [[sku, v["product"], v["variant"], v["vendor"], v["in_stock"], v["quality"], v["size"], v["url"]]
            for sku, v in sorted(site.items(), key=lambda kv: (kv[1]["product"], kv[1]["variant"]))
            if sku in msl_skus]

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    live = sum(1 for r in scrub if r[3] == "Yes")
    run = [["Sub-Category", ", ".join(subcats)], ["Stages", ", ".join(STAGES)], ["Stock buffer", BUFFER],
           ["Run", f"{datetime.datetime.now():%m/%d/%Y %I:%M %p}"], [],
           ["Should be online", len(scrub)], ["Online", live], ["Outliers", len(outliers)], [],
           ["Outliers by reason", ""]]
    for rsn in order + sorted({r[2] for r in outliers} - set(order)):
        n = sum(1 for r in outliers if r[2] == rsn)
        if n:
            run.append([rsn, n])
    run += [[], ["Image quality (online SKUs)", ""]]
    for q in ["Good", "OK", "Low", "No variant photo", "No image"]:
        run.append([q, sum(1 for r in scrub if r[4] == q)])
    ws = write_tab(wb, "Run", ["Website Scrub", ""], run, [32, 40])
    ws.auto_filter.ref = None
    write_tab(wb, "Outliers", ["SKU", "Name", "Reason", "Inventory", "Inventory Source", "Shopify Status",
                               "Sub-Category", "Vendor", "Stage"], outliers, [18, 45, 28, 10, 18, 14, 24, 30, 12])
    write_tab(wb, "Scrub", ["SKU", "Name", "Variant", "Active", "Image Quality", "Inventory", "Inventory Source",
                            "Shopify Status", "Sub-Category", "Vendor", "Stage", "Image Size"],
              scrub, [18, 45, 28, 8, 18, 10, 18, 14, 24, 30, 12, 24])
    ws = write_tab(wb, "HTML", ["SKU", "Product Name", "Variant", "Vendor", "In Stock", "Image Quality",
                                "Image Size", "URL"], html, [18, 45, 28, 30, 9, 18, 24, 60])
    for row in ws.iter_rows(min_row=2):
        row[7].hyperlink = row[7].value

    name = subcats[0] if len(subcats) == 1 else f"{len(subcats)} subcats"
    out = os.path.join(out_dir, f"F&T Site Scrub {name.replace('/', '-')} {datetime.date.today():%m.%d.%y}.xlsx")
    wb.save(out)
    print("\n".join(f"{r[0]}: {r[1]}" for r in run if r))
    print("saved", out)
    return out


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--msl", required=True, help="workbook with the raw MSL export (any sheet with SKU/Stage/MH columns)")
    parser.add_argument("--ds", help="DS inventory report (sheet with SKU + Quantity); optional")
    parser.add_argument("--shopify", required=True, help="shopify_variants.tsv from the admin pull")
    parser.add_argument("--subcat", action="append", required=True, help="MH Sub-Category; repeat for more than one")
    parser.add_argument("--site-json", help="reuse a saved products.json pull instead of fetching")
    parser.add_argument("--out-dir", default=".")
    args = parser.parse_args()
    prods = json.load(open(args.site_json, encoding="utf-8")) if args.site_json else None
    build(args.msl, args.ds, args.shopify, args.subcat, args.out_dir, prods)


if __name__ == "__main__":
    main()
